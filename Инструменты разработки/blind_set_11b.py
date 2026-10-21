#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Задание 11б, шаг 4 (точка Г): короткий слепой набор «Сравнение 3.10» - Ященко:
дообученный F5 / F5 как сейчас / живой чтец (фрагмент его аудиокниги той же длины).

Номера записей - случайные (зерно в ключе), громкость выровнена (ffmpeg loudnorm -18 LUFS),
чтобы не выдавала. Ключ - отдельным файлом в docs («не открывать до оценки»).
Записи синтеза - полным конвейером книги (paragraph_11b.py), живой - кусок его книги.

    python "Инструменты разработки/blind_set_11b.py" <дообученный.mp3> <F5 как сейчас.mp3>
"""
from __future__ import annotations

import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
_os.chdir(_sys.path[0])

import random
import subprocess

DST = "Сравнение 3.10 - лучшие варианты/Слепой набор"
KEY = "docs/ключ 3.10 - слепой набор Ященко (не открывать до оценки).md"
LIVE_SRC = "Образцы голосов/Каменистый Артем - Альфа-ноль 11, Убийца легенд/0003 - Убийца легенд.mp3"
LIVE_AT, LIVE_LEN = 300.0, 60.0


def norm(src: str, dst: str, start: float | None = None, length: float | None = None) -> None:
    cmd = ["ffmpeg", "-v", "error", "-y"]
    if start is not None:
        cmd += ["-ss", str(start), "-t", str(length)]
    cmd += ["-i", src, "-af", "loudnorm=I=-18:TP=-1.5:LRA=11", "-ar", "44100", "-ac", "1", "-b:a", "192k", dst]
    subprocess.run(cmd, check=True)


def sheet(n: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Оценки"
    head = ["Запись", "Похожесть на Ященко (1-5)", "Естественность, манера (1-5)", "Интонация (1-5)",
            "Чистота звука (1-5)", "Комментарий"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor="E8EEF7")
    for k in range(1, n + 1):
        ws.append([f"Запись {k}", None, None, None, None, None])
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        showErrorMessage=True, errorTitle="Оценка", error="Число от 1 до 5")
    ws.add_data_validation(dv)
    dv.add(f"B2:E{n + 1}")
    for col, wdt in zip("ABCDEF", (12, 18, 20, 14, 16, 60)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    wb.save(_os.path.join(DST, "Лист оценок.xlsx"))


def main() -> int:
    trained, current = _sys.argv[1], _sys.argv[2]
    _os.makedirs(DST, exist_ok=True)
    seed = random.SystemRandom().randrange(10 ** 6)
    order = [("дообученный F5 (шаг 4)", trained, None), ("F5 как сейчас", current, None),
             ("живой чтец - его аудиокнига («Альфа-ноль 11», файл 0003, с 5:00, 60 с)", LIVE_SRC, LIVE_AT)]
    random.Random(seed).shuffle(order)
    lines = []
    for k, (what, src, at) in enumerate(order, 1):
        norm(src, _os.path.join(DST, f"Запись {k}.mp3"), at, LIVE_LEN if at is not None else None)
        lines.append(f"| Запись {k} | {what} | `{src}` |")
    sheet(len(order))
    open(_os.path.join(DST, "Как слушать.md"), "w", encoding="utf-8").write(
        "# Как слушать: слепой набор 3.10 (Игорь Ященко)\n\n"
        "Три записи примерно по минуте: две - один и тот же абзац «Главы первой» (синтез, полным "
        "конвейером книги), одна - живой Игорь Ященко, кусок его аудиокниги (другой текст). Что под каким "
        "номером - не сказано, номера случайные, громкость выровнена.\n\n"
        "Оцените в «Лист оценок.xlsx» каждую запись числом 1-5: похожесть на Ященко, естественность и "
        "манера чтения, интонация (не «на одной ноте»), чистота звука (шипение, металл, шум). В "
        "комментарии - что резануло слух (разорванные слова, пропуски, ударения).\n\n"
        "Ключ - в docs: «ключ 3.10 - слепой набор Ященко (не открывать до оценки).md». Открывайте "
        "после того, как заполните лист.\n")
    open(KEY, "w", encoding="utf-8").write(
        "# Ключ: слепой набор 3.10, Ященко (не открывать до оценки)\n\n"
        f"Зерно перемешивания: {seed}. Громкость: ffmpeg loudnorm -18 LUFS. Синтез - полным конвейером "
        "(convert_book), текст - «Глава первая», абзацы 2-4 (`Архив/11б/абзац.txt`).\n\n"
        "| запись | что | исходный файл |\n|---|---|---|\n" + "\n".join(lines) + "\n")
    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
