#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Слепое сравнение движков 3.8 (задание 10а, п. 2; уточнения Андрея 18.10).

Движки: Qwen3-TTS ru-stress 8 бит, он же 4 бита (отдельно), F5 Misha (текущий),
Chatterbox Multilingual V3, VoxCPM2. Голоса - Герасимов и Ященко, те же образцы.
20 фраз (Инструменты разработки/фразы_38.txt): 17 коротких (упрямые слова,
омографы, редкие слова, ударная гласная в начале слова) и 3 длинных абзаца.

Этапы (тяжёлые - под ../.heavy-lock):
    synth  - звук всех движков: Архив/сравнение движков 3.8/<движок>/<голос>__NN.wav
    check  - проверка слов распознаванием (доля ошибок в словах, WER, как в
             Seed-TTS-eval; распознаватель - Whisper large-v3-turbo + нормализация текста) у ВСЕХ движков
             одинаково; пропуск (3+ слова подряд) или WER > 0,15 - пересчёт с другим
             зерном, до двух раз, берётся лучший; итог до и после - проверка.json
    build  - набор для Андрея: «Сравнение 3.8 - движки/Запись NN.mp3» (движок и
             голос - под номером, номера случайные), образцы чтецов, лист оценок
             Excel, инструкция на страницу; ключ - отдельным файлом в docs; в ключ -
             счёт meter2 (ударения) и WER до/после пересчёта
"""
from __future__ import annotations

import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
_os.chdir(_sys.path[0])

import argparse
import difflib
import json
import random
import re
import subprocess

import numpy as np

ENGINES = {"qwen": "Qwen3-TTS 1.7B ru-stress-cf, 8 бит (mlx-audio)",
           "qwen4": "Qwen3-TTS 1.7B ru-stress-cf, 4 бита (mlx-audio)",
           "f5": "F5-TTS Misha - текущий движок Herald («Сбалансировано»)",
           "chatterbox": "Chatterbox Multilingual V3 (mlx-audio)",
           "voxcpm2": "VoxCPM2, 8 бит (mlx-audio, 48 кГц)"}
VOICES = ("Вячеслав Герасимов", "Игорь Ященко")
OUT = "Архив/сравнение движков 3.8"
FOLDER = "Сравнение 3.8 - движки"
KEY = "docs/ключ 3.8 - движки (не открывать до конца прослушивания).md"
SEED = 38
WER_MAX = 0.15
RETRIES = 2


def rows():
    out = []
    for l in open("Инструменты разработки/фразы_38.txt", encoding="utf-8"):
        if l.strip() and not l.startswith("#"):
            ph, w = l.rstrip("\n").rsplit("|", 1)
            out.append((ph, None if w == "-" else w))
    return out


def marked_text(phrase: str, word: str | None) -> str:
    """Фраза с ударениями RUAccent; у слова набора - ударение, как в наборе."""
    import core.pipeline as P
    t = P._with_end_punct(P._prepare_text(phrase, "+", True, True, "ru"))
    if word:
        plain = word.replace("+", "")
        pat = re.compile(r"(?<![\w+])" + r"\+?".join(map(re.escape, plain)) + r"(?![\w])",
                         re.IGNORECASE)
        t = pat.sub(word, t, count=1)
    return t


def make_engine(name):
    if name == "f5":
        from core.engines import get_engine, clone_engine
        from core import profiles
        return get_engine(clone_engine(), language="ru",
                          **profiles.engine_kwargs(profiles.get(None), "balanced"))
    from core.engines.mlxaudio import MlxAudioEngine
    if name in ("qwen", "qwen4"):
        return MlxAudioEngine(model="qwen-stress", bits=4 if name == "qwen4" else 8)
    return MlxAudioEngine(model=name, bits=8)


def wav_path(eng, voice, i):
    return _os.path.join(OUT, eng, f"{voice}__{i + 1:02d}.wav")


def synth_record(eng, name: str, voice: str, i: int, text: str, seed: int) -> np.ndarray:
    """Одна запись (фраза или абзац); длинный абзац - кусками, как в программе."""
    import core.pipeline as P
    pieces = (P._chunk_paragraphs(text, eng.max_chunk_chars, 0.4, 0.8, 1.1)
              if len(text) > eng.max_chunk_chars else [(text, 0.0, 1.0, ())])
    parts = []
    for k, pc in enumerate(pieces):
        txt = P._with_end_punct(pc[0])
        if name == "f5":
            import mlx.core as mx
            mx.random.seed(seed * 10 + k)
            au = eng.synth_batch([txt], voice)[0]
        else:
            au = eng.synth_chunk(txt, voice, seed=seed * 10 + k)
        parts += [np.asarray(au, np.float32), np.zeros(int(pc[1] * eng.sample_rate), np.float32)]
    return np.concatenate(parts)


def synth(a) -> int:
    import soundfile as sf
    R = rows()
    texts = [marked_text(p, w) for p, w in R]
    for name in a.engines.split(","):
        _os.makedirs(_os.path.join(OUT, name), exist_ok=True)
        todo = [(v, i) for v in VOICES for i in range(len(texts))
                if not _os.path.exists(wav_path(name, v, i))]
        if not todo:
            continue
        eng = make_engine(name)
        eng.ensure_loaded()
        for voice, i in todo:
            au = synth_record(eng, name, voice, i, texts[i], SEED + i)
            sf.write(wav_path(name, voice, i), au, eng.sample_rate)
            print(f"{name} / {voice}: {i + 1}/{len(texts)}", flush=True)
        if hasattr(eng, "close"):
            eng.close()
        del eng
    return 0


# --- проверка слов распознаванием ---------------------------------------------------
def _words(t: str) -> list[str]:
    from core.normalize import normalize_text
    t = normalize_text(t.replace("+", ""), expand_numbers=True, language="ru")
    return re.findall(r"[а-яё]+", t.lower().replace("ё", "е"))


def wer_of(audio, sr, ref: str) -> tuple[float, list[str]]:
    """(доля ошибок в словах, пропуски - 3+ слова подряд)."""
    import core.pipeline as P
    h = P._listen(audio, sr, "ru", P._ASR_FULL)
    hyp = _words(" ".join(w for w, _s, _e in (h or [])))
    want = _words(ref)
    sm = difflib.SequenceMatcher(None, want, hyp, autojunk=False)
    err = 0
    skips = []
    for t, a1, a2, b1, b2 in sm.get_opcodes():
        if t == "equal":
            continue
        err += max(a2 - a1, b2 - b1)
        if t in ("delete", "replace") and a2 - a1 >= 3 and (b2 - b1) <= (a2 - a1) // 2:
            skips.append(" ".join(want[a1:a2]))
    return err / max(1, len(want)), skips


def check(a) -> int:
    import soundfile as sf
    R = rows()
    texts = [marked_text(p, w) for p, w in R]
    res_path = _os.path.join(OUT, "проверка.json")
    res = json.load(open(res_path, encoding="utf-8")) if _os.path.exists(res_path) else {}
    for name in a.engines.split(","):
        eng = None
        for voice in VOICES:
            for i in range(len(texts)):
                k = f"{name}|{voice}|{i + 1}"
                if k in res and "after" in res[k]:
                    continue
                au, sr = sf.read(wav_path(name, voice, i), dtype="float32")
                w0, s0 = (1.0, ["тишина"]) if len(au) / sr < 0.5 else wer_of(au, sr, R[i][0])
                best = (len(s0), w0, au, s0)
                tries = 0
                while (best[0] or best[1] > WER_MAX) and tries < RETRIES:
                    tries += 1
                    if eng is None:
                        eng = make_engine(name)
                        eng.ensure_loaded()
                    nu = synth_record(eng, name, voice, i, texts[i], SEED + 1000 * tries + i)
                    sr = eng.sample_rate
                    w1, s1 = (1.0, ["тишина"]) if len(nu) / sr < 0.5 else wer_of(nu, sr, R[i][0])
                    if (len(s1), w1) < best[:2]:
                        best = (len(s1), w1, nu, s1)
                        sf.write(wav_path(name, voice, i), nu, sr)
                res[k] = {"before": {"wer": round(w0, 3), "skips": s0},
                          "after": {"wer": round(best[1], 3), "skips": best[3]}, "retries": tries}
                json.dump(res, open(res_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            print(f"{name} / {voice}: проверено", flush=True)
        if eng is not None and hasattr(eng, "close"):
            eng.close()
        del eng
    print(table(res))
    return 0


def table(res: dict) -> str:
    """Таблица: по движку - средний WER и записи с пропусками, до и после пересчёта."""
    lines = ["| движок | WER до | WER после | записей с пропусками до | после | пересчитано записей |",
             "|---|---|---|---|---|---|"]
    for name in ENGINES:
        rs = [v for k, v in res.items() if k.startswith(name + "|")]
        if not rs:
            continue
        n = len(rs)
        lines.append(f"| {ENGINES[name]} | {sum(r['before']['wer'] for r in rs) / n:.1%} | "
                     f"{sum(r['after']['wer'] for r in rs) / n:.1%} | "
                     f"{sum(bool(r['before']['skips']) for r in rs)} из {n} | "
                     f"{sum(bool(r['after']['skips']) for r in rs)} из {n} | "
                     f"{sum(r['retries'] > 0 for r in rs)} |")
    return "\n".join(lines)


# --- набор для Андрея -----------------------------------------------------------------
def _level(a: np.ndarray) -> np.ndarray:
    """Одна громкость у всех записей: RMS речи -> 0,07, пики не выше 0,95."""
    a = np.asarray(a, np.float32)
    voiced = a[np.abs(a) > 0.01]
    rms = float(np.sqrt(np.mean(voiced ** 2))) if len(voiced) else 0.0
    if rms > 0:
        a = a * (0.07 / rms)
    pk = float(np.abs(a).max()) if len(a) else 0.0
    return a * (0.95 / pk) if pk > 0.95 else a


def build(a) -> int:
    import soundfile as sf
    from core import meter
    from core.stress import to_combining
    from core.engines.f5 import default_voices_dir
    R = rows()
    stress_idx = [i for i, (_p, w) in enumerate(R) if w]
    names = [n for n in ENGINES if all(_os.path.exists(wav_path(n, v, i))
                                       for v in VOICES for i in range(len(R)))]
    # meter2: ударение в слове набора
    auto = {}
    for name in names:
        for voice in VOICES:
            ok = 0
            wrong = []
            for i in stress_idx:
                au, sr = sf.read(wav_path(name, voice, i), dtype="float32")
                ph, w = R[i]
                bad = meter.wrong_words(au, sr, marked_text(ph, w),
                                        only={w.replace("+", "").lower()}, min_margin=0.0)
                ok += not bad
                if bad:
                    wrong.append(w)
            auto[(name, voice)] = (ok, len(stress_idx), wrong)
            print(f"meter2: {name} / {voice}: верно {ok} из {len(stress_idx)}", flush=True)
    # записи под случайными номерами
    recs = [(n, v) for n in names for v in VOICES]
    rnd = random.Random("Herald 3.8, слепое сравнение движков")
    rnd.shuffle(recs)
    import shutil
    shutil.rmtree(FOLDER, ignore_errors=True)
    _os.makedirs(FOLDER, exist_ok=True)
    for num, (name, voice) in enumerate(recs, 1):
        parts, sr0 = [], 48000
        for i in range(len(R)):
            au, sr = sf.read(wav_path(name, voice, i), dtype="float32")
            if sr != sr0:
                from math import gcd
                from scipy.signal import resample_poly
                g = gcd(sr, sr0)
                au = resample_poly(au, sr0 // g, sr // g).astype(np.float32)
            parts += [_level(au), np.zeros(int(1.5 * sr0), np.float32)]
        tmp = _os.path.join(OUT, f"запись {num:02d}.wav")
        sf.write(tmp, np.concatenate(parts), sr0)
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", tmp, "-b:a", "256k",
                        _os.path.join(FOLDER, f"Запись {num:02d}.mp3")], check=True)
    # образцы чтецов - для «похожести голоса» (только в папке для прослушивания, не в git)
    for voice in VOICES:
        src = _os.path.join(default_voices_dir(), voice + ".wav")
        if _os.path.exists(src):
            subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", src, "-b:a", "256k",
                            _os.path.join(FOLDER, f"Образец чтеца - {voice}.mp3")], check=True)
    sheet(recs, R)
    instruction(len(recs), R)
    # ключ
    try:
        res = json.load(open(_os.path.join(OUT, "проверка.json"), encoding="utf-8"))
    except (OSError, ValueError):
        res = {}
    key = ["# Ключ 3.8 - слепое сравнение движков", "",
           "Открывать после того, как лист оценок заполнен.", "",
           "| запись | движок | чтец | ударения по meter2 (17 слов) | WER до → после пересчёта |",
           "|---|---|---|---|---|"]
    for num, (name, voice) in enumerate(recs, 1):
        ok, cnt, wrong = auto[(name, voice)]
        rs = [v for k, v in res.items() if k.startswith(f"{name}|{voice}|")]
        wer = (f"{sum(r['before']['wer'] for r in rs) / len(rs):.1%} → "
               f"{sum(r['after']['wer'] for r in rs) / len(rs):.1%}") if rs else "-"
        key.append(f"| {num:02d} | {ENGINES[name]} | {voice} | {ok} из {cnt}"
                   + (f" (не так: {', '.join(to_combining(w.lower()) for w in wrong)})" if wrong else "")
                   + f" | {wer} |")
    key += ["", "Проверка слов распознаванием (доля ошибок в словах, как в Seed-TTS-eval; распознаватель - Whisper large-v3-turbo), у всех "
            "движков одинаково: пропуск 3+ слов или WER > 15 % - пересчёт с другим зерном, до двух "
            "раз, взят лучший.", "", table(res) if res else ""]
    open(KEY, "w", encoding="utf-8").write("\n".join(key) + "\n")
    json.dump({f"{n}|{v}": list(x) for (n, v), x in auto.items()},
              open(_os.path.join(OUT, "meter2.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return 0


def sheet(recs, R) -> None:
    """Лист оценок Excel: номер записи, чтец; естественность, ударения, похожесть
    голоса - 1-5; комментарий. Второй лист - фразы и верные ударения."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from core.stress import to_combining
    wb = Workbook()
    ws = wb.active
    ws.title = "Оценки"
    head = ["Запись", "Чтец", "Естественность (1-5)", "Ударения (1-5)",
            "Похожесть голоса (1-5)", "Комментарий"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor="E8EEF7")
    for num, (_name, voice) in enumerate(recs, 1):
        ws.append([f"Запись {num:02d}", voice, None, None, None, None])
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        showErrorMessage=True, errorTitle="Оценка", error="Число от 1 до 5")
    ws.add_data_validation(dv)
    dv.add(f"C2:E{len(recs) + 1}")
    for col, wdt in zip("ABCDEF", (12, 22, 16, 14, 18, 60)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    w2 = wb.create_sheet("Фразы")
    w2.append(["№", "Фраза (в каждой записи - по порядку)", "Слово с верным ударением"])
    for c in w2[1]:
        c.font = Font(bold=True)
    for i, (ph, w) in enumerate(R, 1):
        w2.append([i, ph, to_combining(w.lower()) if w else "- (длинный абзац: звук)"])
    w2.column_dimensions["A"].width = 5
    w2.column_dimensions["B"].width = 90
    w2.column_dimensions["C"].width = 26
    for row in w2.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(_os.path.join(FOLDER, "Лист оценок.xlsx"))


def instruction(n: int, R) -> None:
    text = f"""# Как слушать: слепое сравнение движков Herald 3.8

**Что это.** {n} записей - одни и те же 20 фраз разными движками синтеза, двумя
чтецами (Вячеслав Герасимов и Игорь Ященко). Какой движок под каким номером - не
сказано, номера случайные. Громкость у всех записей выровнена.

**В каждой записи по порядку** (пауза 1,5 с между фразами):
- фразы 1-17 - короткие, в каждой одно трудное слово (список и верные ударения -
  лист «Фразы» в Excel): «о́тмель», «ка́торжник», «ча́ек», «долги́», «по́черком»,
  за́мок / замо́к, се́ла / села́, редкие слова, слова с ударной гласной в начале;
- фразы 18-20 - три длинных абзаца из «Проб пера» - для оценки звука.

**Что оценить** - в «Лист оценок.xlsx», лист «Оценки», по каждой записи числом 1-5:
- **Естественность** - как живой чтец: интонация, паузы, нет металла, шипения, «робота».
- **Ударения** - в трудных словах 1-17 (5 - все верно, 1 - ошибок много).
- **Похожесть голоса** - насколько похоже на чтеца. Для сравнения в папке лежат
  «Образец чтеца - Вячеслав Герасимов.mp3» и «… - Игорь Ященко.mp3».
- **Комментарий** - что резануло слух: пропущенные слова, обрывы, странные места.

**Как удобнее.** Достаточно прослушать фразы 1-17 и один абзац; если запись
явно хуже других - можно не дослушивать и поставить оценки.

**Ключ** (какой движок под каким номером) - отдельным файлом в docs:
«ключ 3.8 - движки (не открывать до конца прослушивания).md». Откройте его
только после того, как лист заполнен.

Решение, какой движок станет основным в Herald, - за вами; до вашей оценки
основным остаётся прежний.
"""
    open(_os.path.join(FOLDER, "Как слушать.md"), "w", encoding="utf-8").write(text)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("stage", choices=["synth", "check", "build", "table"])
    ap.add_argument("--engines", default=",".join(ENGINES))
    a = ap.parse_args()
    if a.stage == "table":
        print(table(json.load(open(_os.path.join(OUT, "проверка.json"), encoding="utf-8"))))
        return 0
    return {"synth": synth, "check": check, "build": build}[a.stage](a)


if __name__ == "__main__":
    rc = main()
    _sys.stdout.flush()
    _os._exit(rc)
