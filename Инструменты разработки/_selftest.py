#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Быстрая проверка ридеров и нарезки на синтетических файлах (без синтеза)."""
import os as _os, sys as _sys
# скрипт лежит в подпапке, а core/ — в корне проекта
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
_os.chdir(_sys.path[0])
import os, tempfile, zipfile
from core.readers import read_document
from core.chapters import build_chapters
from core.normalize import normalize_text

D = tempfile.mkdtemp(prefix="t2a_test_")

# --- txt с заголовками, числами, прямой речью ---
txt = """Предисловие к тесту, чтобы было что читать до первой главы. Тут просто текст.

Глава 1. Начало
В 1990 г. случилось важное. До поезда оставалось 45 минут, а решение так и не пришло.
- Это ты, Пётр? - тихо спросила она.

Глава 2. Развязка
Стоимость составила 1 250 000 руб., т.е. заметно дороже. См. стр. 12.
"""
p_txt = os.path.join(D, "sample.txt"); open(p_txt, "w").write(txt)

# --- docx со стилями заголовков ---
import docx
d = docx.Document()
d.core_properties.title = "Тестовая книга"
d.core_properties.author = "Автор Тестов"
d.add_heading("Глава первая", level=1)
d.add_paragraph("Текст первой главы. Было 3,5 килограмма и 12 процентов.")
d.add_heading("Глава вторая", level=1)
d.add_paragraph("Второй кусок. § 5 и № 7.")
p_docx = os.path.join(D, "sample.docx"); d.save(p_docx)

# --- fb2 (минимальный) ---
fb2 = """<?xml version="1.0" encoding="utf-8"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" xmlns:l="http://www.w3.org/1999/xlink">
<description><title-info>
<book-title>Книга ФБ2</book-title>
<author><first-name>Иван</first-name><last-name>Петров</last-name></author>
</title-info></description>
<body>
<section><title><p>Глава один</p></title><p>Первый абзац фб2.</p><p>Второй абзац.</p></section>
<section><title><p>Глава два</p></title><p>Текст второй главы.</p></section>
</body></FictionBook>"""
p_fb2 = os.path.join(D, "sample.fb2"); open(p_fb2, "w").write(fb2)

for p in (p_txt, p_docx, p_fb2):
    doc = read_document(p)
    chs = build_chapters(doc)
    print(f"\n### {os.path.basename(p)}  title={doc.title!r} author={doc.author!r} cover={bool(doc.cover)}")
    for c in chs:
        print(f"   - {c.title!r}  ({len(c.text)} симв.)")

print("\n### нормализация:")
print("  ", normalize_text("В 1990 г. было 1 250 000 руб., т.е. дорого. § 5, № 7, 3,5 кг, 12%."))
print("\nOK, tmpdir:", D)
