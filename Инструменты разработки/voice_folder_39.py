#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Задание 11: папка «Сравнение 3.9 - голоса диктора на выбор» (по мере готовности).

На голос - один файл «Голос NN.mp3»: кусок с трудными ударениями (о́тмель,
о́тмели, ка́торжник, ча́ек) и абзац из книги; имя без названия движка. Лист
оценок - короткий. Какой голос откуда - Архив/голоса диктора 3.9/папка.json
(для отчёта, не для Андрея).

    python "Инструменты разработки/voice_folder_39.py" "native (куски)|Голос 03" "icl_fixed (куски)|Игорь Ященко" ...
Новые голоса дописываются следующими номерами; уже положенные не меняются.
"""
from __future__ import annotations

import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
_os.chdir(_sys.path[0])

import importlib.util
import json
import subprocess

import numpy as np

BANK = "Архив/голоса диктора 3.9"
FOLDER = "Сравнение 3.9 - голоса диктора на выбор"
PIECES = (0, 5)                       # кусок 1 (4 трудных слова) и абзац 1


def main() -> int:
    import soundfile as sf
    spec = importlib.util.spec_from_file_location("e", "Инструменты разработки/bench_engines_38.py")
    e = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(e)
    _os.makedirs(FOLDER, exist_ok=True)
    mp = _os.path.join(BANK, "папка.json")
    mapping = json.load(open(mp, encoding="utf-8")) if _os.path.exists(mp) else []
    for spec_ in _sys.argv[1:]:
        if any(m["источник"] == spec_ for m in mapping):
            continue
        mode, label = spec_.split("|", 1)
        parts = []
        for k in PIECES:
            au, sr = sf.read(_os.path.join(BANK, mode, f"{label}__{k + 1:02d}.wav"), dtype="float32")
            parts += [e._level(au), np.zeros(int(1.2 * sr), np.float32)]
        num = len(mapping) + 1
        tmp = _os.path.join(BANK, f"папка {num:02d}.wav")
        sf.write(tmp, np.concatenate(parts), sr)
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", tmp, "-b:a", "256k",
                        _os.path.join(FOLDER, f"Голос {num:02d}.mp3")], check=True)
        _os.remove(tmp)
        mapping.append({"файл": f"Голос {num:02d}", "источник": spec_})
        print(f"Голос {num:02d} <- {spec_}", flush=True)
    json.dump(mapping, open(mp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    sheet(len(mapping))
    note(len(mapping))
    return 0


def sheet(n: int) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    path = _os.path.join(FOLDER, "Лист оценок.xlsx")
    if _os.path.exists(path):                      # оценки Андрея не трогаем - только дописываем строки
        wb = load_workbook(path)
        ws = wb["Оценки"]
        have = ws.max_row - 1
        for k in range(have + 1, n + 1):
            ws.append([f"Голос {k:02d}", None, None, None, None])
        wb.save(path)
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Оценки"
    ws.append(["Голос", "Нравится голос (1-5)", "Ударения (1-5)", "Качество звука (1-5)", "Комментарий"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor="E8EEF7")
    for k in range(1, n + 1):
        ws.append([f"Голос {k:02d}", None, None, None, None])
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        showErrorMessage=True, errorTitle="Оценка", error="Число от 1 до 5")
    ws.add_data_validation(dv)
    dv.add("B2:D60")
    for col, w in zip("ABCDE", (12, 18, 14, 18, 60)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def note(n: int) -> None:
    text = f"""# Голоса диктора на выбор (Herald 3.9)

Сейчас голосов: {n}. Папка пополняется по мере готовности - новые голоса
появляются следующими номерами, уже лежащие не меняются.

В каждом файле: четыре фразы с трудными ударениями («Северная о́тмель»,
«С о́тмели…», «ка́торжник», «крик ча́ек») и абзац из «Пробы пера».

Оцените в «Лист оценок.xlsx» (1-5): нравится ли голос для книг, верны ли
ударения, качество звука (шипение, металл, ровность); комментарий - по желанию.
Выбранный голос станет голосом Herald после вашего решения - по умолчанию ничего
не переключено.
"""
    open(_os.path.join(FOLDER, "Как выбирать.md"), "w", encoding="utf-8").write(text)


if __name__ == "__main__":
    rc = main()
    _sys.stdout.flush()
    _os._exit(rc)
