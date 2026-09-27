"""Чтение входных форматов в единый объект Document.

Поддержка: .txt/.md, .docx, .epub, .fb2 (+ .fb2.zip), .pdf (только текстовый слой).
PDF-сканы без текстового слоя намеренно отсеиваются с понятной ошибкой — OCR не
делаем (об этом договорились).

Каждый ридер возвращает Document с метаданными, обложкой (если есть в файле) и
либо готовыми главами, либо сплошным текстом — тогда главы построит chapters.py.
"""

from __future__ import annotations

import os
import re
import zipfile
from dataclasses import dataclass, field
from typing import Optional


class ReaderError(Exception):
    pass


class ScannedPdfError(ReaderError):
    """PDF без текстового слоя (скан). Такой файл пропускаем."""


@dataclass
class CoverImage:
    data: bytes
    ext: str = "jpg"          # jpg / png


@dataclass
class Chapter:
    title: str
    text: str


@dataclass
class Document:
    source_path: str
    title: str = ""
    author: str = ""
    cover: Optional[CoverImage] = None
    chapters: Optional[list[Chapter]] = None      # если структура известна
    text: Optional[str] = None                    # если структуры нет

    @property
    def base_title(self) -> str:
        return self.title or os.path.splitext(os.path.basename(self.source_path))[0]


SUPPORTED = {".txt", ".md", ".docx", ".epub", ".fb2", ".pdf", ".zip"}


def read_document(path: str) -> Document:
    ext = os.path.splitext(path)[1].lower()
    low = path.lower()
    if low.endswith(".fb2.zip"):
        return _read_fb2(path, zipped=True)
    if ext in (".txt", ".md"):
        return _read_txt(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".epub":
        return _read_epub(path)
    if ext == ".fb2":
        return _read_fb2(path)
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".zip":
        # zip с одним fb2 внутри — частый случай библиотек
        return _read_fb2(path, zipped=True)
    raise ReaderError(f"Формат не поддерживается: {ext}")


# --- txt / md ------------------------------------------------------------
def _read_txt(path: str) -> Document:
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "utf-8", "cp1251", "koi8-r", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    return Document(source_path=path, text=text)


# --- docx ----------------------------------------------------------------
_HEADING_RE = re.compile(r"(heading|заголовок)\s*([1-3])", re.IGNORECASE)


def _read_docx(path: str) -> Document:
    import docx
    d = docx.Document(path)
    title = (d.core_properties.title or "").strip()
    author = (d.core_properties.author or "").strip()

    chapters: list[Chapter] = []
    cur_title, cur_buf = None, []

    def flush():
        if cur_buf and any(p.strip() for p in cur_buf):
            chapters.append(Chapter(cur_title or "", "\n\n".join(cur_buf).strip()))

    for p in d.paragraphs:
        style = (p.style.name if p.style else "") or ""
        txt = p.text.strip()
        if not txt:
            continue
        if style.lower() in ("title", "заголовок") and not title:
            title = txt
            continue
        if _HEADING_RE.search(style):
            flush()
            cur_title, cur_buf = txt, []
        else:
            cur_buf.append(txt)
    flush()

    if chapters:
        return Document(source_path=path, title=title, author=author, chapters=chapters)
    # заголовков нет — отдаём сплошным текстом
    full = "\n\n".join(p.text.strip() for p in d.paragraphs if p.text.strip())
    return Document(source_path=path, title=title, author=author, text=full)


# --- epub ----------------------------------------------------------------
def _read_epub(path: str) -> Document:
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup

    book = epub.read_epub(path)

    def meta(ns, name):
        try:
            m = book.get_metadata(ns, name)
            return m[0][0].strip() if m else ""
        except Exception:
            return ""

    title = meta("DC", "title")
    author = meta("DC", "creator")

    # карта href -> заголовок из оглавления
    toc_titles: dict[str, str] = {}

    def walk_toc(items):
        for it in items:
            if isinstance(it, tuple):
                section, children = it
                if getattr(section, "href", None):
                    toc_titles[section.href.split("#")[0]] = section.title
                walk_toc(children)
            else:
                if getattr(it, "href", None):
                    toc_titles[it.href.split("#")[0]] = it.title

    try:
        walk_toc(book.toc)
    except Exception:
        pass

    chapters: list[Chapter] = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        name = item.get_name()
        soup = BeautifulSoup(item.get_content(), "lxml")
        # заголовок: из оглавления, иначе первый h1..h3
        ctitle = toc_titles.get(name, "")
        if not ctitle:
            h = soup.find(["h1", "h2", "h3"])
            ctitle = h.get_text(" ", strip=True) if h else ""
        text = soup.get_text("\n", strip=True)
        if len(text) < 30:          # обложки, служебные страницы
            continue
        chapters.append(Chapter(ctitle, text))

    cover = _epub_cover(book)
    if not chapters:
        return Document(source_path=path, title=title, author=author, cover=cover, text="")
    return Document(source_path=path, title=title, author=author, cover=cover, chapters=chapters)


def _epub_cover(book) -> Optional[CoverImage]:
    import ebooklib
    # 1) явный тип COVER
    try:
        for it in book.get_items_of_type(ebooklib.ITEM_COVER):
            if it.get_content():
                return _cover_from_bytes(it.get_content(), it.get_name())
    except Exception:
        pass
    # 2) meta name=cover -> id изображения
    try:
        covers = book.get_metadata("OPF", "cover")
        if covers:
            cid = covers[0][1].get("content")
            it = book.get_item_with_id(cid)
            if it and it.get_content():
                return _cover_from_bytes(it.get_content(), it.get_name())
    except Exception:
        pass
    # 3) первое изображение с "cover" в имени
    try:
        for it in book.get_items_of_type(ebooklib.ITEM_IMAGE):
            if "cover" in it.get_name().lower():
                return _cover_from_bytes(it.get_content(), it.get_name())
    except Exception:
        pass
    return None


def _cover_from_bytes(data: bytes, name: str) -> CoverImage:
    ext = "png" if name.lower().endswith("png") else "jpg"
    return CoverImage(data=data, ext=ext)


# --- fb2 -----------------------------------------------------------------
def _read_fb2(path: str, zipped: bool = False) -> Document:
    from lxml import etree
    import base64

    if zipped:
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if n.lower().endswith(".fb2")] or z.namelist()
            data = z.read(names[0])
    else:
        data = open(path, "rb").read()

    parser = etree.XMLParser(recover=True, huge_tree=True)
    root = etree.fromstring(data, parser=parser)

    def L(node, name):
        """Namespace-agnostic поиск по локальному имени."""
        return node.xpath(f".//*[local-name()='{name}']")

    def text_of(node) -> str:
        parts = []
        for p in L(node, "p"):
            t = "".join(p.itertext()).strip()
            if t:
                parts.append(t)
        return "\n\n".join(parts)

    title_info = L(root, "title-info")
    ti = title_info[0] if title_info else root
    bt = L(ti, "book-title")
    title = bt[0].text.strip() if bt and bt[0].text else ""
    authors = []
    for a in L(ti, "author"):
        fn = L(a, "first-name")
        ln = L(a, "last-name")
        name = " ".join(x[0].text.strip() for x in (fn, ln) if x and x[0].text)
        if name:
            authors.append(name)
    author = ", ".join(authors)

    # обложка: coverpage -> image href="#id" -> binary id
    cover = None
    cp = L(ti, "coverpage")
    if cp:
        imgs = L(cp[0], "image")
        href = ""
        if imgs:
            for k, v in imgs[0].attrib.items():
                if k.endswith("href"):
                    href = v.lstrip("#")
        if href:
            for b in L(root, "binary"):
                if b.get("id") == href and b.text:
                    ctype = b.get("content-type", "image/jpeg")
                    ext = "png" if "png" in ctype else "jpg"
                    try:
                        cover = CoverImage(base64.b64decode(b.text), ext)
                    except Exception:
                        cover = None
                    break

    bodies = L(root, "body")
    chapters: list[Chapter] = []
    if bodies:
        sections = L(bodies[0], "section")
        # берём только секции верхнего уровня (не вложенные)
        top = [s for s in sections if s.getparent() is bodies[0]]
        targets = top or sections
        for s in targets:
            tnodes = L(s, "title")
            ctitle = ""
            if tnodes:
                ctitle = " ".join("".join(tnodes[0].itertext()).split())
            body_text = text_of(s)
            if body_text.strip():
                chapters.append(Chapter(ctitle, body_text))

    if chapters:
        return Document(source_path=path, title=title, author=author, cover=cover, chapters=chapters)
    full = text_of(bodies[0]) if bodies else ""
    return Document(source_path=path, title=title, author=author, cover=cover, text=full)


# --- pdf (только текстовый слой) -----------------------------------------
def _read_pdf(path: str) -> Document:
    import pymupdf
    doc = pymupdf.open(path)
    n = doc.page_count
    pages = [doc.load_page(i).get_text("text") for i in range(n)]
    total = sum(len(p.strip()) for p in pages)
    # эвристика скана: почти нет извлекаемого текста
    if total < max(200, 25 * n):
        doc.close()
        raise ScannedPdfError(
            f"В PDF нет текстового слоя (похоже на скан): {os.path.basename(path)}. "
            f"Такие файлы пропускаем — нужен OCR, а мы его не делаем."
        )
    md = doc.metadata or {}
    title = (md.get("title") or "").strip()
    author = (md.get("author") or "").strip()
    doc.close()
    return Document(source_path=path, title=title, author=author, text="\n\n".join(pages))
